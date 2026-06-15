"""Excel report exporter for IP3 using openpyxl."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="E3ECFB")
_HEADER_FONT = Font(bold=True, color="1D4ED8")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _autowidth(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export(case: dict, results: list[dict], out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Sheet 1: idea + conditions
    ws = wb.active
    ws.title = "아이디어"
    ws.append(["항목", "내용"])
    _style_header(ws, 2)
    ws.append(["아이디어명", case.get("title", "")])
    ws.append(["아이디어 설명", case.get("description", "")])
    ws.append(["검색 범위", case.get("search_scope", "")])
    ws.append(["결과 수", case.get("top_n", "")])
    ws.append(["핵심 키워드", case.get("keywords", "")])
    ws.append(["제외어", case.get("exclude_keywords", "")])
    structure = case.get("idea_structure") or {}
    if isinstance(structure, dict):
        for k, v in structure.items():
            ws.append([f"핵심 구성 - {k}", v])
    _autowidth(ws, [22, 90])
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = _WRAP

    # Sheet 2: comparison summary
    ws2 = wb.create_sheet("비교분석 요약")
    ws2.append(["특허명", "주요 유사점", "주요 차이점", "확인할 점", "판단 상태"])
    _style_header(ws2, 5)
    for r in results:
        ws2.append([
            r.get("title", ""),
            "; ".join(r.get("similar_points", [])),
            "; ".join(r.get("different_points", [])),
            "; ".join(r.get("check_points", [])),
            r.get("judgment_status", ""),
        ])
    _autowidth(ws2, [40, 45, 45, 40, 12])
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP

    # Sheet 3: per-patent detail
    ws3 = wb.create_sheet("특허별 상세")
    ws3.append([
        "특허명", "국가", "출원/공개번호", "유사한 점", "차이점",
        "확인할 점", "확인 위치", "원문 근거", "판단 상태", "검토 메모",
    ])
    _style_header(ws3, 10)
    for r in results:
        ws3.append([
            r.get("title", ""),
            r.get("country", ""),
            r.get("publication_no", "") or r.get("application_no", ""),
            "\n".join(r.get("similar_points", [])),
            "\n".join(r.get("different_points", [])),
            "\n".join(r.get("check_points", [])),
            ", ".join(r.get("source_locations", [])),
            "\n".join(r.get("original_evidence", [])),
            r.get("judgment_status", ""),
            r.get("summary_memo", ""),
        ])
    _autowidth(ws3, [30, 8, 18, 35, 35, 35, 18, 35, 12, 30])
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP

    wb.save(out_path)
    return out_path
